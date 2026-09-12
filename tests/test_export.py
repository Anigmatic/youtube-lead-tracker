import csv

import openpyxl

from app.export import export_csv, export_xlsx

_LEADS = [
    {
        "date": "2026-09-11", "language": "English", "name": "Test Channel",
        "channel_url": "https://www.youtube.com/@testchannel",
        "subscriber_count_display": "5.65K", "avg_views_display": "1.5K-4.1K",
        "contact_info": "contact@testchannel.com", "fit_assessment": "High",
        "fit_reason": "Matches niche.", "status": "New", "outreach_method": "Email", "notes": "",
    },
]


def test_export_csv_writes_header_and_rows(tmp_path):
    path = tmp_path / "leads.csv"
    export_csv(_LEADS, str(path))
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0][0] == "Date"
    assert "Test Channel" in rows[1]
    assert "contact@testchannel.com" in rows[1]


def test_export_xlsx_writes_header_and_rows(tmp_path):
    path = tmp_path / "leads.xlsx"
    export_xlsx(_LEADS, str(path))
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    data_row = [cell.value for cell in ws[2]]
    assert header[0] == "Date"
    assert "Test Channel" in data_row


def test_export_csv_includes_links_column(tmp_path):
    leads = [dict(_LEADS[0], links="site.com, twitter.com/test, patreon.com/test")]
    path = tmp_path / "leads.csv"
    export_csv(leads, str(path))
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert "Links" in rows[0]
    assert "site.com, twitter.com/test, patreon.com/test" in rows[1]


_FORMULA_LEADS = [
    {
        "date": "2026-09-11", "language": "English",
        "name": '=HYPERLINK("http://evil.test","Claim prize")',
        "channel_url": "https://www.youtube.com/@evil",
        "subscriber_count_display": "5.65K", "avg_views_display": "1.5K-4.1K",
        "contact_info": "contact@testchannel.com", "fit_assessment": "High",
        "fit_reason": "Matches niche.", "status": "New", "outreach_method": "Email", "notes": "",
    },
]


def test_export_csv_escapes_formula_injection_in_name(tmp_path):
    path = tmp_path / "leads.csv"
    export_csv(_FORMULA_LEADS, str(path))
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    name_cell = rows[1][2]
    assert name_cell.startswith("'=")
    assert not name_cell.startswith("=")


def test_export_xlsx_escapes_formula_injection_in_name(tmp_path):
    path = tmp_path / "leads.xlsx"
    export_xlsx(_FORMULA_LEADS, str(path))
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    name_cell = ws[2][2].value
    assert name_cell.startswith("'=")
    assert not name_cell.startswith("=")
